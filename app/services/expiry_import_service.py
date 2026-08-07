"""Импорт Excel-базы сроков годности и остатков в SQLite.

Как и с базой сертификатов, Excel — единственный источник данных, SQLite —
только его быстрый индекс, и при каждом импорте таблица переписывается
целиком.

Структура листа фиксирована требованием продукта: колонка A — полное
системное название товара, B — срок годности, C — количество (остаток).
Один товар может встречаться несколько раз — по одной строке на партию.
"""

from __future__ import annotations

import logging
import time
from pathlib import Path
from typing import Any

from app.core.dates import parse_date
from app.core.events import AppEvent, EventBus
from app.core.exceptions import ExcelParsingError
from app.models.expiry_stock import ExpiryStockEntry
from app.models.history import ImportResult
from app.repositories.expiry_stock_repository import ExpiryStockRepository
from app.repositories.settings_repository import (
    KEY_EXPIRY_STOCK_EXCEL_FINGERPRINT,
    KEY_LAST_EXPIRY_STOCK_EXCEL_PATH,
    KEY_LAST_EXPIRY_STOCK_IMPORT_AT,
    SettingsRepository,
)

_LOGGER = logging.getLogger(__name__)

#: Позиции колонок (0-based) в листе базы сроков годности.
_COLUMN_PRODUCT_NAME = 0
_COLUMN_EXPIRY = 1
_COLUMN_QUANTITY = 2


class ExpiryImportService:
    """Читает Excel-базу сроков годности и переносит её в SQLite."""

    def __init__(
        self,
        repository: ExpiryStockRepository,
        settings_repository: SettingsRepository,
        event_bus: EventBus,
    ) -> None:
        """Создаёт сервис импорта.

        Args:
            repository: Репозиторий Excel-базы сроков годности.
            settings_repository: Служебное состояние (путь и отпечаток файла).
            event_bus: Шина событий приложения.
        """
        self._repository = repository
        self._state = settings_repository
        self._events = event_bus

    def import_file(self, path: Path) -> ImportResult:
        """Читает файл базы сроков годности и полностью переимпортирует её.

        Args:
            path: Путь к Excel-файлу базы сроков годности.

        Returns:
            Итог импорта.
        """
        started = time.perf_counter()
        self._events.publish(AppEvent.EXPIRY_STOCK_IMPORT_STARTED, path=str(path))

        try:
            entries, skipped = self._read_file(path)
        except ExcelParsingError as error:
            result = ImportResult(
                duration_ms=(time.perf_counter() - started) * 1000, error=error.message
            )
            _LOGGER.error("Импорт базы сроков годности не выполнен: %s", error.message)
            self._events.publish(AppEvent.EXPIRY_STOCK_IMPORT_FINISHED, result=result)
            return result

        imported = self._repository.replace_all(entries)

        self._state.set(KEY_LAST_EXPIRY_STOCK_EXCEL_PATH, str(path))
        self._state.set(KEY_EXPIRY_STOCK_EXCEL_FINGERPRINT, _fingerprint(path))
        self._state.set_now(KEY_LAST_EXPIRY_STOCK_IMPORT_AT)

        result = ImportResult(
            total_rows=imported + skipped,
            imported=imported,
            skipped_rows=skipped,
            duration_ms=(time.perf_counter() - started) * 1000,
        )
        _LOGGER.info(
            "%s за %.2f с", result.summary("База сроков годности"), result.duration_ms / 1000
        )
        self._events.publish(AppEvent.EXPIRY_STOCK_IMPORT_FINISHED, result=result)
        return result

    # ------------------------------------------------------------- состояние
    def last_path(self) -> str:
        """Возвращает путь последнего импортированного файла.

        Returns:
            Путь либо пустая строка.
        """
        return self._state.get(KEY_LAST_EXPIRY_STOCK_EXCEL_PATH)

    def last_imported_at(self) -> str:
        """Возвращает время последнего импорта в виде строки.

        Returns:
            Отметка времени либо пустая строка.
        """
        value = self._state.get_datetime(KEY_LAST_EXPIRY_STOCK_IMPORT_AT)
        return value.strftime("%d.%m.%Y %H:%M") if value else ""

    def needs_reimport(self, path: Path) -> bool:
        """Проверяет, изменился ли файл базы со времени последнего импорта.

        Args:
            path: Путь к файлу базы сроков годности.

        Returns:
            ``True``, если файл нужно переимпортировать.
        """
        if not path.is_file():
            return False
        if str(path) != self.last_path():
            return True
        return _fingerprint(path) != self._state.get(KEY_EXPIRY_STOCK_EXCEL_FINGERPRINT)

    # ------------------------------------------------------------- чтение
    def _read_file(self, path: Path) -> tuple[list[ExpiryStockEntry], int]:
        """Читает лист базы сроков годности и строит объекты.

        Args:
            path: Путь к Excel-файлу.

        Returns:
            Пара «прочитанные партии, количество пропущенных строк».

        Raises:
            ExcelParsingError: Если файл не удалось открыть.
        """
        try:
            from openpyxl import load_workbook
        except ImportError as error:  # pragma: no cover — зависимость обязательна
            raise ExcelParsingError("Библиотека openpyxl недоступна", path=path) from error

        if not path.exists():
            raise ExcelParsingError(f"Файл не найден: {path}", path=path)

        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as error:
            raise ExcelParsingError(
                f"Не удалось открыть файл базы сроков годности: {path.name}",
                path=path,
                details=str(error),
            ) from error

        try:
            sheet = workbook.active
            if sheet is None:
                raise ExcelParsingError(f"В файле {path.name} нет листов", path=path)

            entries: list[ExpiryStockEntry] = []
            skipped = 0
            for index, values in enumerate(sheet.iter_rows(values_only=True)):
                if index == 0 and _looks_like_header(values):
                    continue
                entry = _row_to_entry(values, source_row=index + 1)
                if entry is None:
                    if any(value is not None and str(value).strip() for value in values):
                        skipped += 1
                    continue
                entries.append(entry)
            return (entries, skipped)
        finally:
            workbook.close()


def _looks_like_header(values: tuple[Any, ...]) -> bool:
    """Проверяет, похожа ли первая строка листа на заголовок таблицы.

    В отличие от базы сертификатов (там колонка A — числовой порядковый
    номер, легко отличимый от текста заголовка), колонка A здесь — название
    товара, и оно текстовое что в заголовке, что в данных. Заголовок
    отличают по колонке B: в данных там всегда срок годности (дата или
    Excel-serial), а в заголовке — текст вроде «Срок годности», который
    датой не распознаётся.
    """
    second = values[_COLUMN_EXPIRY] if len(values) > _COLUMN_EXPIRY else None
    return parse_date(second) is None


def _row_to_entry(values: tuple[Any, ...], *, source_row: int) -> ExpiryStockEntry | None:
    """Преобразует строку листа в :class:`ExpiryStockEntry`.

    Args:
        values: Значения ячеек строки.
        source_row: Порядковый номер строки в исходном файле (1-based).

    Returns:
        Партия товара либо ``None``, если наименование (обязательное поле)
        в строке отсутствует.
    """
    product_name = _cell_text(values, _COLUMN_PRODUCT_NAME)
    if not product_name:
        return None

    expiry_raw = values[_COLUMN_EXPIRY] if len(values) > _COLUMN_EXPIRY else None
    quantity_raw = values[_COLUMN_QUANTITY] if len(values) > _COLUMN_QUANTITY else None

    return ExpiryStockEntry(
        product_name=product_name,
        expiry_date=parse_date(expiry_raw),
        quantity=_to_float(quantity_raw),
        source_row=source_row,
    )


def _cell_text(values: tuple[Any, ...], index: int) -> str:
    """Возвращает текстовое значение ячейки, безопасно обрабатывая границы."""
    if index >= len(values) or values[index] is None:
        return ""
    text = str(values[index]).strip()
    return "" if text.lower() in {"none", "nan", "-", "—"} else text


def _to_float(value: Any) -> float:
    """Приводит значение ячейки количества к ``float``, ``0.0`` при неудаче."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", ".").replace(" ", ""))
    except ValueError:
        return 0.0


def _fingerprint(path: Path) -> str:
    """Строит лёгкий отпечаток файла (размер + время изменения)."""
    try:
        stat = path.stat()
    except OSError:
        return ""
    return f"{stat.st_size}:{stat.st_mtime}"
