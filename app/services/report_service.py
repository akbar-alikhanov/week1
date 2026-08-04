"""Формирование отчёта о результатах подбора.

Отчёт — рабочий документ для проверки: он показывает, какой сертификат
подставлен в каждую строку, с какой уверенностью и каким методом. Именно по
нему специалист проверяет спорные позиции, не открывая исходный заказ.

Данные готовятся через pandas: это даёт бесплатную сводную статистику и
единый код для выгрузки в XLSX и CSV.
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any

from app.core.constants import APP_COMPANY, APP_NAME, APP_VERSION
from app.core.exceptions import ExportError
from app.models.history import ProcessingRun
from app.models.matching import MatchResult, MatchStatus

_LOGGER = logging.getLogger(__name__)


class ReportService:
    """Строит отчёт о совпадениях в формате XLSX или CSV."""

    def build_dataframe(self, results: list[MatchResult]) -> Any:
        """Формирует таблицу результатов.

        Args:
            results: Результаты подбора.

        Returns:
            :class:`pandas.DataFrame` с одной строкой на позицию заказа.

        Raises:
            ExportError: Если pandas недоступен.
        """
        try:
            import pandas as pd
        except ImportError as error:  # pragma: no cover — зависимость обязательна
            raise ExportError("Библиотека pandas недоступна", details=str(error)) from error

        records = [result.to_report_dict() for result in results]
        return pd.DataFrame(records)

    def build_summary(self, results: list[MatchResult], run: ProcessingRun | None = None) -> Any:
        """Формирует сводку по обработке.

        Args:
            results: Результаты подбора.
            run: Сведения о запуске (для длительности и путей).

        Returns:
            :class:`pandas.DataFrame` с показателями «параметр — значение».
        """
        import pandas as pd

        counters: dict[str, int] = {}
        for result in results:
            counters[result.status.label] = counters.get(result.status.label, 0) + 1

        total = len(results)
        resolved = sum(1 for result in results if result.status.is_resolved)
        scores = [result.score for result in results if result.score > 0]

        rows: list[tuple[str, Any]] = [
            ("Программа", f"{APP_NAME} {APP_VERSION}"),
            ("Организация", APP_COMPANY),
            ("Дата обработки", datetime.now().strftime("%d.%m.%Y %H:%M")),
            ("Файл заказа", run.order_name if run else ""),
            ("Всего позиций", total),
            ("Подобрано автоматически", counters.get(MatchStatus.MATCHED.label, 0)),
            ("Выбрано вручную", counters.get(MatchStatus.MANUAL.label, 0)),
            ("Требует проверки", counters.get(MatchStatus.AMBIGUOUS.label, 0)),
            ("Не найдено", counters.get(MatchStatus.NOT_FOUND.label, 0)),
            ("Доля успешных, %", round(resolved / total * 100, 1) if total else 0.0),
            (
                "Средняя уверенность, %",
                round(sum(scores) / len(scores) * 100, 1) if scores else 0.0,
            ),
        ]
        if run is not None:
            rows.append(("Длительность, с", round(run.duration_seconds, 2)))
            rows.append(("Скопировано файлов", run.copied_files))
        return pd.DataFrame(rows, columns=["Показатель", "Значение"])

    def write(
        self,
        results: list[MatchResult],
        destination: Path,
        *,
        run: ProcessingRun | None = None,
        report_format: str = "xlsx",
    ) -> Path:
        """Сохраняет отчёт в файл.

        Args:
            results: Результаты подбора.
            destination: Путь файла отчёта (расширение будет приведено к формату).
            run: Сведения о запуске.
            report_format: ``xlsx`` или ``csv``.

        Returns:
            Путь созданного файла.

        Raises:
            ExportError: Если файл не удалось записать.
        """
        frame = self.build_dataframe(results)
        destination.parent.mkdir(parents=True, exist_ok=True)

        if report_format.lower() == "csv":
            target = destination.with_suffix(".csv")
            try:
                frame.to_csv(target, index=False, sep=";", encoding="utf-8-sig")
            except OSError as error:
                raise ExportError("Не удалось сохранить отчёт", details=str(error)) from error
            _LOGGER.info("Отчёт сохранён: %s", target)
            return target

        target = destination.with_suffix(".xlsx")
        summary = self.build_summary(results, run)
        problems = (
            frame[frame["Статус"].isin([MatchStatus.AMBIGUOUS.label, MatchStatus.NOT_FOUND.label])]
            if not frame.empty
            else frame
        )

        try:
            with self._writer(target) as writer:
                summary.to_excel(writer, sheet_name="Сводка", index=False)
                frame.to_excel(writer, sheet_name="Результаты", index=False)
                if not problems.empty:
                    problems.to_excel(writer, sheet_name="Требует внимания", index=False)
                self._style(writer, frame, summary, problems)
        except ExportError:
            raise
        except Exception as error:
            raise ExportError("Не удалось сохранить отчёт", details=str(error)) from error

        _LOGGER.info("Отчёт сохранён: %s", target)
        return target

    def _writer(self, target: Path) -> Any:
        """Создаёт writer pandas для записи книги Excel."""
        import pandas as pd

        return pd.ExcelWriter(target, engine="openpyxl")

    def _style(self, writer: Any, frame: Any, summary: Any, problems: Any) -> None:
        """Применяет оформление к листам отчёта."""
        try:
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:  # pragma: no cover
            return

        header_fill = PatternFill("solid", fgColor="2F5D62")
        header_font = Font(bold=True, color="FFFFFF")
        widths = {
            "Сводка": [28, 34],
            "Результаты": [8, 46, 14, 18, 20, 18, 40, 22, 14, 20, 34, 12, 34, 12],
            "Требует внимания": [8, 46, 14, 18, 20, 18, 40, 22, 14, 20, 34, 12, 34, 12],
        }

        for sheet_name, sheet in writer.sheets.items():
            column_count = sheet.max_column
            for column in range(1, column_count + 1):
                cell = sheet.cell(row=1, column=column)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for index, width in enumerate(widths.get(sheet_name, []), start=1):
                if index <= column_count:
                    sheet.column_dimensions[get_column_letter(index)].width = width
            sheet.freeze_panes = "A2"
            if sheet.max_row > 1 and sheet_name != "Сводка":
                sheet.auto_filter.ref = f"A1:{get_column_letter(column_count)}{sheet.max_row}"

        _ = (frame, summary, problems)  # данные уже записаны, ссылки для читаемости
